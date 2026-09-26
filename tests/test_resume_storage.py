from io import BytesIO

from docx import Document

from web.resume_storage import extract_resume_text, read_resume


def test_docx_resume_text_includes_paragraphs_and_tables():
    document = Document()
    document.add_paragraph("Example Person")
    document.add_paragraph("Customer service and front desk experience")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Employer"
    table.cell(0, 1).text = "Example Hotel"
    output = BytesIO()
    document.save(output)

    text = extract_resume_text("resume.docx", output.getvalue())

    assert "Example Person" in text
    assert "front desk experience" in text
    assert "Example Hotel" in text


def test_read_resume_reads_local_storage_uri(tmp_path):
    path = tmp_path / "resume.pdf"
    path.write_bytes(b"resume bytes")

    assert read_resume(str(path)) == b"resume bytes"
