\
from __future__ import annotations

import os
from io import BytesIO
from pathlib import Path
from datetime import datetime, timezone

def save_resume(
    filename: str,
    content: bytes,
    resume_type: str,
    user_id: str | None = None,
) -> str:
    bucket_name = os.getenv("GCS_BUCKET")
    safe_name = Path(filename).name
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    owner = user_id or "legacy"
    object_name = f"resumes/{owner}/{resume_type}/{stamp}-{safe_name}"

    if bucket_name:
        from google.cloud import storage
        client = storage.Client()
        bucket = client.bucket(bucket_name)
        blob = bucket.blob(object_name)
        blob.upload_from_string(content)
        return f"gs://{bucket_name}/{object_name}"

    upload_dir = (
        Path(__file__).resolve().parents[1]
        / "data" / "uploads" / owner / resume_type
    )
    upload_dir.mkdir(parents=True, exist_ok=True)
    path = upload_dir / f"{stamp}-{safe_name}"
    path.write_bytes(content)
    return str(path)


def read_resume(storage_uri: str) -> bytes:
    if storage_uri.startswith("gs://"):
        bucket_name, object_name = storage_uri[5:].split("/", 1)
        from google.cloud import storage
        client = storage.Client()
        return client.bucket(bucket_name).blob(object_name).download_as_bytes()
    return Path(storage_uri).read_bytes()


def extract_resume_text(filename: str, content: bytes) -> str:
    lower_name = filename.lower()
    if lower_name.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in reader.pages).strip()
    if lower_name.endswith(".docx"):
        from docx import Document
        document = Document(BytesIO(content))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(value for value in paragraphs if value.strip()).strip()
    raise ValueError("Only DOCX and PDF resumes are supported.")



def delete_resume(storage_uri: str) -> None:
    if storage_uri.startswith("gs://"):
        bucket_name, object_name = storage_uri[5:].split("/", 1)
        from google.cloud import storage
        client = storage.Client()
        client.bucket(bucket_name).blob(object_name).delete()
        return
    path = Path(storage_uri)
    if path.exists() and path.is_file():
        path.unlink()
